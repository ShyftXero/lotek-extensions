"""Per-report section order: the persisted ReportBoard.section_order + layouts.resolve_section_order — the
ONE resolver both renderers loop, so the HTML preview and the DOCX/PDF deliverable agree on section order.

The resolver is deliberately forgiving of anything persisted or POSTed: unknown keys dropped, dups
collapsed, missing blocks appended disabled, None => default. These are its contract, pinned here.
"""
from __future__ import annotations

import uuid

import scribble.models as fm
from scribble.reporting.layouts import (
    BLOCK_KEYS,
    DEFAULT_SECTION_ORDER,
    SectionSpec,
    default_section_specs,
    enabled_section_keys,
    resolve_section_order,
)


def test_none_and_empty_resolve_to_the_default_composition():
    for raw in (None, [], {}, "nonsense"):
        specs = resolve_section_order(raw)
        assert [s.key for s in specs] == list(DEFAULT_SECTION_ORDER)
        # every block present exactly once, all enabled except the opt-in activity_log
        assert {s.key for s in specs} == set(BLOCK_KEYS)
        assert [s.key for s in specs if not s.enabled] == ["activity_log"]


def test_default_specs_match_the_default_resolution():
    assert default_section_specs() == resolve_section_order(None)


def test_saved_order_is_respected_with_enabled_flags():
    raw = [
        {"key": "cover", "enabled": True},
        {"key": "summary", "enabled": True},
        {"key": "methodology", "enabled": True},
        {"key": "findings", "enabled": True},
        {"key": "evidence", "enabled": False},  # explicitly turned off
    ]
    specs = resolve_section_order(raw)
    # the five saved come first, in saved order
    assert [s.key for s in specs[:5]] == ["cover", "summary", "methodology", "findings", "evidence"]
    assert specs[4] == SectionSpec("evidence", False)
    # enabled_section_keys drops the disabled one AND every unmentioned (appended-disabled) block
    assert enabled_section_keys(raw) == ("cover", "summary", "methodology", "findings")


def test_missing_blocks_are_appended_disabled_not_silently_added():
    # A report saved with only two sections must not gain the rest as visible sections.
    raw = [{"key": "cover", "enabled": True}, {"key": "findings", "enabled": True}]
    specs = resolve_section_order(raw)
    assert {s.key for s in specs} == set(BLOCK_KEYS)  # still lists the whole vocabulary (for the editor)
    appended = [s for s in specs if s.key not in ("cover", "findings")]
    assert all(not s.enabled for s in appended)  # ...but every appended block is OFF
    assert enabled_section_keys(raw) == ("cover", "findings")


def test_unknown_keys_dropped_and_dups_collapse_to_first():
    raw = [
        {"key": "summary", "enabled": True},
        {"key": "bogus", "enabled": True},          # unknown -> dropped
        {"key": "summary", "enabled": False},       # dup -> collapses to the first (enabled)
        "findings",                                  # bare string -> enabled
    ]
    specs = resolve_section_order(raw)
    assert "bogus" not in {s.key for s in specs}
    summary = [s for s in specs if s.key == "summary"]
    assert summary == [SectionSpec("summary", True)]  # first wins, not the later disabled dup
    assert SectionSpec("findings", True) in specs


def test_all_disabled_yields_an_empty_render_order():
    raw = [{"key": k, "enabled": False} for k in BLOCK_KEYS]
    assert enabled_section_keys(raw) == ()


def test_section_spec_label_is_human_readable():
    assert SectionSpec("summary", True).label == "Executive Summary"
    assert SectionSpec("evidence", False).label == "Evidence Appendix"


def test_section_order_column_round_trips(session_factory):
    saved = [{"key": "cover", "enabled": True}, {"key": "findings", "enabled": True},
             {"key": "summary", "enabled": False}]
    with session_factory() as db:
        eng = fm.ReportBoard(name="Order Eng", client_id=uuid.uuid7(), section_order=saved)
        db.add(eng)
        db.commit()
        eng_id = eng.id
    with session_factory() as db:
        got = db.get(fm.ReportBoard, eng_id)
        assert got.section_order == saved
        assert enabled_section_keys(got.section_order) == ("cover", "findings")


def test_section_order_defaults_null(session_factory):
    with session_factory() as db:
        eng = fm.ReportBoard(name="Default Eng", client_id=uuid.uuid7())
        db.add(eng)
        db.commit()
        eng_id = eng.id
    with session_factory() as db:
        got = db.get(fm.ReportBoard, eng_id)
        assert got.section_order is None  # NULL => default composition at resolve time
        assert enabled_section_keys(got.section_order)[0] == "cover"


# --- the HTML renderer honors the persisted per-report order (Task #2) ---------------------------------

def _render_html_for(session_factory, section_order, **kw) -> str:
    from scribble.content import schema
    from scribble.enums import Severity
    from scribble.reporting import build_report_context
    from scribble.reporting.render_html import render_report_html

    with session_factory() as db:
        eng = fm.ReportBoard(name="HTML Order Eng", client_id=uuid.uuid7(), company_name="Acme",
                             section_order=section_order)
        grp = fm.FindingGroup(engagement=eng, name="Internal", order_index=0)
        db.add_all([eng, grp])
        db.flush()
        db.add(fm.BoardFinding(engagement_id=eng.id, group_id=grp.id, order_index=0, title="V",
                               severity=Severity.high,
                               content_json={"description": schema.doc_from_text("x")}))
        db.commit()
        eng_id = eng.id
    with session_factory() as db:
        return render_report_html(build_report_context(db.get(fm.ReportBoard, eng_id)), **kw)


def test_html_honors_persisted_order_and_drops_unlisted(session_factory):
    # findings BEFORE summary (opposite of the default preset) proves the persisted order drives the HTML;
    # methodology is absent from the saved list -> appended disabled -> its section is dropped even though it
    # renders by default (standing text). Anchors are id="sec-<key>".
    order = [{"key": "findings", "enabled": True}, {"key": "summary", "enabled": True}]
    html = _render_html_for(session_factory, order)
    assert html.index('id="sec-findings"') < html.index('id="sec-summary"')
    assert 'id="sec-methodology"' not in html


def test_explicit_layout_query_previews_a_preset_over_the_persisted_order(session_factory):
    # A saved order puts findings before summary; an explicit ?layout=default PREVIEWS the standard preset
    # (summary before findings, methodology present) for that render, without touching what's persisted.
    order = [{"key": "findings", "enabled": True}, {"key": "summary", "enabled": True}]
    html = _render_html_for(session_factory, order, layout="default")
    assert html.index('id="sec-summary"') < html.index('id="sec-findings"')
    assert 'id="sec-methodology"' in html  # the preset carries it; the persisted order had dropped it


# --- the DOCX deliverable honors the persisted per-report order (Task #3, marker reorder) --------------

def _docx_headings(payload: bytes) -> list[str]:
    import io

    import docx
    d = docx.Document(io.BytesIO(payload))
    return [p.text for p in d.paragraphs
            if p.style and p.style.name.startswith("Heading") and (p.text or "").strip()]


def _render_docx_for(session_factory, section_order) -> bytes:
    from scribble.content import schema
    from scribble.enums import Severity
    from scribble.reporting import build_report_context
    from scribble.reporting.render_docx import render_report_docx

    with session_factory() as db:
        eng = fm.ReportBoard(name="DOCX Order Eng", client_id=uuid.uuid7(), company_name="Acme",
                             section_order=section_order)
        grp = fm.FindingGroup(engagement=eng, name="Internal", order_index=0)
        db.add_all([eng, grp])
        db.flush()
        db.add(fm.BoardFinding(engagement_id=eng.id, group_id=grp.id, order_index=0, title="V",
                               severity=Severity.high,
                               content_json={"description": schema.doc_from_text("x")}))
        db.commit()
        eng_id = eng.id
    with session_factory() as db:
        return render_report_docx(build_report_context(db.get(fm.ReportBoard, eng_id)))


def test_docx_honors_persisted_order_methodology_before_findings(session_factory):
    order = [{"key": "summary", "enabled": True}, {"key": "methodology", "enabled": True},
             {"key": "findings", "enabled": True}]
    headings = _docx_headings(_render_docx_for(session_factory, order))
    assert "Methodology" in headings and "Findings" in headings
    assert headings.index("Methodology") < headings.index("Findings")


def test_docx_drops_sections_absent_from_the_saved_order(session_factory):
    # only summary + findings saved -> methodology + TOC are appended-disabled -> dropped from the DOCX.
    order = [{"key": "summary", "enabled": True}, {"key": "findings", "enabled": True}]
    headings = _docx_headings(_render_docx_for(session_factory, order))
    assert "Findings" in headings
    assert "Methodology" not in headings
    assert "Table of Contents" not in headings


def test_docx_default_order_is_the_standard_sequence(session_factory):
    # NULL section_order -> default composition -> the standard sequence. Regression guard for the reorder:
    # no leftover section-marker bookmarks, standard order preserved.
    import io

    import docx
    from docx.oxml.ns import qn
    payload = _render_docx_for(session_factory, None)
    headings = _docx_headings(payload)
    assert headings.index("Executive Summary") < headings.index("Findings") < headings.index("Methodology")
    assert "Table of Contents" in headings
    d = docx.Document(io.BytesIO(payload))
    leftover = [bm.get(qn("w:name")) for bm in d.element.body.iter(qn("w:bookmarkStart"))
                if (bm.get(qn("w:name")) or "").startswith("scribble-section:")]
    assert leftover == [], "section-marker bookmarks must be removed by the reorder"
